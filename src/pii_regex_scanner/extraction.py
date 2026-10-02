"""Stage 1 extraction helpers: file detection, decoding, and structured parsing.

This module is where raw files become scanable source units. It contains the
logic for:

- recognising file types from content as well as filename;
- hashing and decoding text safely;
- extracting bounded records/rows from structured and tabular sources;
- routing documents and images through the appropriate parser path.

The implementation is loaded into the shared pipeline namespace so the package
and standalone forms use the same extraction behaviour.
"""

def has_email_evidence(evidence_types):
    """Return whether a set of evidence types contains any email category."""
    return bool(set(evidence_types) & EMAIL_EVIDENCE_TYPES)


ID_EVIDENCE_TYPES = set()
SCAN_NLP_ENABLED = False
SCAN_NLP_BACKEND = ""


load_rules(DEFAULT_RULES_PATH)


def has_id_evidence(evidence_types):
    """Return whether a set of evidence types contains any ID-like evidence."""
    return bool(set(evidence_types) & ID_EVIDENCE_TYPES)


def set_nlp_mode(value):
    """Enable or disable NLP-assisted header/key resolution for scanning."""
    global SCAN_NLP_ENABLED, SCAN_NLP_BACKEND
    SCAN_NLP_ENABLED = bool(value)
    if not SCAN_NLP_ENABLED:
        SCAN_NLP_BACKEND = ""
    TABLE_HEADER_MATCH_CACHE.clear()
    TABLE_HEADER_PATTERN_CACHE.clear()


def current_nlp_backend():
    return SCAN_NLP_BACKEND if SCAN_NLP_ENABLED else ""


def has_risk_role(evidence_types, role):
    """Return whether any evidence type in the set carries the requested role."""
    return any(role in EVIDENCE_RISK_ROLES.get(evidence_type, set()) for evidence_type in evidence_types)


def hash_file(path, algorithm="sha256", progress_callback=None, total_size=None, progress_base=0.0, progress_span=100.0):
    """Hash one file with the requested algorithm, optionally emitting progress."""
    algorithm = safe_cell(algorithm or "sha256").casefold()
    if algorithm not in {"sha256", "md5"}:
        raise ValueError(f"Unsupported file hash algorithm: {algorithm}")
    h = hashlib.new(algorithm)
    bytes_done = 0
    last_progress = 0.0
    with open(path, "rb", buffering=HASH_CHUNK_SIZE) as f:
        for chunk in iter(lambda: f.read(HASH_CHUNK_SIZE), b""):
            h.update(chunk)
            bytes_done += len(chunk)
            if progress_callback and total_size and total_size > 0:
                now = time.monotonic()
                if now - last_progress >= 1.0:
                    percent = progress_base + progress_span * min(1.0, bytes_done / total_size)
                    progress_callback(percent, "hashing")
                    last_progress = now
    if progress_callback:
        progress_callback(progress_base + progress_span, "hashing")
    return h.hexdigest()


def md5_file(path, progress_callback=None, total_size=None, progress_base=0.0, progress_span=100.0):
    """Compatibility wrapper for callers/tests that still request MD5."""
    return hash_file(
        path,
        "md5",
        progress_callback=progress_callback,
        total_size=total_size,
        progress_base=progress_base,
        progress_span=progress_span,
    )


def normalize_value(evidence_type, value):
    """Normalize a matched value according to its evidence type."""
    value = value.strip()
    if evidence_type in {"Citizenship Country", "Citizenship/Country"}:
        normalized_country = normalize_citizenship_country_value(value)
        return normalized_country or value
    normalization = EVIDENCE_NORMALIZATION.get(evidence_type, "none")

    if normalization == "lower":
        return value.lower()
    if normalization == "compact_upper":
        return re.sub(r"[\s\-]", "", value).upper()
    if normalization == "address_compact":
        return normalize_address_value(value)
    if normalization == "name_casefold":
        return normalize_name_value(value)
    if normalization == "country_name":
        return normalize_citizenship_country_value(value)
    if normalization not in {"", "none"}:
        raise ValueError(f"Unknown normalization strategy for {evidence_type}: {normalization}")
    return value


def clean_context(text, start, end, window=80):
    """Return a compact context window around a match for review output."""
    left = max(0, start - window)
    right = min(len(text), end + window)
    context = text[left:right]
    context = context.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    context = re.sub(r"\s+", " ", context).strip()
    return context


def extract_match_value(match):
    """Return the value-bearing part of a regex match.

    Rules can expose a named `value` group when the full regex contains leading
    labels or context that should not become the emitted matched value.
    """
    groupdict = match.groupdict()
    if "value" in groupdict and groupdict["value"]:
        return groupdict["value"], match.start("value"), match.end("value")
    return match.group(0), match.start(), match.end()


def is_binary_like(raw):
    """Heuristically detect binary content that should not be decoded as text."""
    if not raw:
        return False
    if looks_like_utf16_text(raw):
        return False
    if b"\x00" in raw[:4096]:
        return True
    control_count = sum(1 for b in raw[:4096] if b < 32 and b not in (9, 10, 13))
    return control_count > max(32, len(raw[:4096]) // 10)


def looks_like_utf16_text(raw):
    """Heuristic UTF-16 detector for files without a reliable BOM."""
    sample = raw[:4096]
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return True
    if len(sample) < 8:
        return False

    even_nulls = sample[0::2].count(0)
    odd_nulls = sample[1::2].count(0)
    half_length = max(1, len(sample) // 2)
    return even_nulls / half_length > 0.25 or odd_nulls / half_length > 0.25


def decode_text_bytes(raw):
    """Decode bytes using the scanner's permissive text-decoding strategy."""
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="ignore")
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig", errors="ignore")
    if looks_like_utf16_text(raw):
        sample = raw[:4096]
        even_nulls = sample[0::2].count(0)
        odd_nulls = sample[1::2].count(0)
        encoding = "utf-16-be" if even_nulls > odd_nulls else "utf-16-le"
        return raw.decode(encoding, errors="ignore")
    return raw.decode("utf-8", errors="ignore")


def decode_person_table_bytes(raw):
    """Decode person-table CSV bytes with broader export-oriented fallbacks."""
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")) or looks_like_utf16_text(raw):
        return decode_text_bytes(raw)

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue

    return raw.decode("utf-8", errors="replace")


def text_like_bytes(raw):
    """Return whether the given bytes are plausible text for scanning."""
    return bool(raw) and not is_binary_like(raw)


def sniff_zip_office_extension(path):
    """Inspect a ZIP container and infer the Office/OpenDocument file family."""
    import zipfile

    if not zipfile.is_zipfile(path):
        return ""

    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        lower_names = {name.lower() for name in names}

        if any(name.startswith("xl/") for name in lower_names):
            return ".xlsx"
        if any(name.startswith("word/") for name in lower_names):
            return ".docx"
        if any(name.startswith("ppt/") for name in lower_names):
            return ".pptx"

        try:
            mimetype = archive.read("mimetype").decode("ascii", errors="ignore").strip()
        except Exception:
            mimetype = ""
        if mimetype == "application/vnd.oasis.opendocument.spreadsheet":
            return ".ods"
        if mimetype == "application/vnd.oasis.opendocument.text":
            return ".odt"

    return ""


def sniff_delimited_text_extension(text):
    """Guess whether text is better treated as CSV or TSV."""
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=TABLE_SNIFF_DELIMITERS)
    except csv.Error:
        return ""
    if dialect.delimiter == "\t":
        return ".tsv"
    return ".csv"


def detect_extension_from_content(path, original_extension):
    """Return the extension to use for extraction based on file signatures.

    Strong binary and markup signatures override the filename extension. If no
    reliable signature is found, known extensions are trusted. Unknown text-like
    files are routed as plain text or delimited tables.
    """
    try:
        with open(path, "rb") as f:
            raw = f.read(8192)
    except OSError:
        return original_extension

    if raw.startswith(b"%PDF"):
        return ".pdf"
    if raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        if original_extension in {".xls", ".doc", ".msg"}:
            return original_extension
        return ".xls"
    if raw.startswith(b"PK\x03\x04") or raw.startswith(b"PK\x05\x06") or raw.startswith(b"PK\x07\x08"):
        sniffed = sniff_zip_office_extension(path)
        if sniffed:
            return sniffed

    stripped = raw.lstrip()
    stripped_lower = stripped[:512].lower()
    if stripped.startswith(b"{\\rtf"):
        return ".rtf"
    if stripped_lower.startswith((b"<!doctype html", b"<html")) or b"<html" in stripped_lower:
        return ".html"
    if stripped_lower.startswith(b"<?xml"):
        return ".xml"

    if raw.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if raw.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if raw.startswith((b"GIF87a", b"GIF89a")):
        return ".gif"
    if raw.startswith((b"II*\x00", b"MM\x00*")):
        return ".tif"
    if raw.startswith(b"BM"):
        return ".bmp"
    if raw.startswith(b"RIFF") and raw[8:12] == b"WEBP":
        return ".webp"
    if raw[4:8] == b"ftyp" and raw[8:12].lower() in {b"heic", b"heix", b"hevc", b"hevx", b"mif1", b"msf1"}:
        return ".heic"

    if original_extension in SCAN_EXTRACTABLE_EXTENSIONS or original_extension in LEGACY_UNSUPPORTED_DOCUMENT_EXTENSIONS:
        return original_extension

    if text_like_bytes(raw):
        text = decode_text_bytes(raw)
        delimited_extension = sniff_delimited_text_extension(text)
        if delimited_extension:
            return delimited_extension
        return ".txt"

    return original_extension


def read_text_sample(path, max_text_chars):
    """Read and decode a bounded text sample from a file for plain-text scanning."""
    chunks = []
    total = 0

    with open(path, "rb") as f:
        while True:
            if max_text_chars > 0:
                raw_limit = max_text_chars * 4
                if total >= raw_limit:
                    break
                read_size = min(READ_CHUNK_SIZE, raw_limit - total)
            else:
                read_size = READ_CHUNK_SIZE

            chunk = f.read(read_size)
            if not chunk:
                break
            chunks.append(chunk)
            total += len(chunk)

    raw = b"".join(chunks)
    if is_binary_like(raw):
        return "", "binary_like_content"

    text = decode_text_bytes(raw)
    if max_text_chars > 0 and len(text) > max_text_chars:
        return text[:max_text_chars], "truncated"
    return text, "scanned"


def limit_extracted_text(extracted, max_text_chars):
    if extracted.extraction_method.startswith("truncated_"):
        return extracted.text, extracted.extraction_method
    if max_text_chars > 0 and len(extracted.text) > max_text_chars:
        return extracted.text[:max_text_chars], f"truncated_{extracted.extraction_method}"
    return extracted.text, f"scanned_{extracted.extraction_method}"




def safe_cell(value):
    """Return a compact single-line representation of a table cell."""
    if value is None:
        return ""
    value = str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", value).strip()


def looks_like_table_headers(headers, header_separator_chars=""):
    cleaned = [safe_cell(h) for h in headers]
    non_empty = [h for h in cleaned if h]
    if len(non_empty) < TABLE_COLUMN_RULES.get("min_non_empty_headers", 2):
        return False
    if TABLE_COLUMN_RULES.get("header_matchers"):
        return any(resolve_table_header(header, header_separator_chars=header_separator_chars) for header in non_empty)
    joined = " | ".join(non_empty).lower()
    return any(term in joined for term in TABLE_COLUMN_RULES.get("header_terms", ()))


def resolve_table_header(header, header_separator_chars=""):
    match = resolve_table_header_match(header, header_separator_chars=header_separator_chars)
    return match["canonical_header"] if match else ""


def resolve_table_header_metadata(header, header_separator_chars=""):
    match = resolve_table_header_match(header, header_separator_chars=header_separator_chars)
    return dict(match["metadata"]) if match else {}


def structured_header_leaf_candidates(header):
    header = safe_cell(header)
    if not header or not re.search(r"[./]", header):
        return []
    if header_label_looks_like_free_text(header):
        return []
    parts = [safe_cell(part) for part in re.split(r"[./]", header) if safe_cell(part)]
    if not parts:
        return []
    candidates = []
    leaf = parts[-1]
    if leaf.casefold() in {"cdata", "text", "value"} and len(parts) >= 2:
        leaf = parts[-2]
    if leaf and leaf != header:
        candidates.append(leaf)
    if len(parts) >= 2:
        pair = ".".join(parts[-2:])
        if pair != header and pair not in candidates:
            candidates.append(pair)
    return candidates


def table_header_match_supported(header, canonical_header, metadata):
    if header_label_looks_like_free_text(header):
        return False
    if canonical_header in {"first name", "last name", "full name"} and is_non_person_name_header(header):
        return False
    guard = metadata.get("context_guard") if isinstance(metadata, dict) else {}
    if guard and not header_context_supports_guard(header, guard):
        return False
    return True


def spaced_header_variant(header, header_separator_chars=""):
    return re.sub(header_word_separator_pattern(header_separator_chars), " ", safe_cell(header)).strip()


def split_header_with_separator(header, separator):
    header = safe_cell(header)
    separator = safe_cell(separator)
    if not header or not separator or separator not in header:
        return ""
    parts = [safe_cell(part) for part in re.split(re.escape(separator) + r"+", header) if safe_cell(part)]
    if len(parts) < 2:
        return ""
    return " ".join(parts)


def resolve_table_header_match_uncached(header, header_separator_chars=""):
    matches = []
    spaced_header = spaced_header_variant(header, header_separator_chars)
    camel_spaced_header = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", spaced_header)
    header_variants = (header, spaced_header, camel_spaced_header)
    for variant_order, header_variant in enumerate(dict.fromkeys(header_variants)):
        for pattern, canonical_header, order, metadata in TABLE_COLUMN_RULES.get("header_matchers", ()):
            match = pattern.search(header_variant)
            if match:
                matches.append(
                    (variant_order, match.start(), -(match.end() - match.start()), order, canonical_header, metadata)
                )
    matches = [
        match
        for match in matches
        if table_header_match_supported(header, match[4], match[5])
    ]
    if not matches:
        return None
    match = min(matches)
    return {
        "canonical_header": match[4],
        "metadata": match[5],
    }


def header_separator_split_supported(header, separator):
    """Return whether splitting one label on a separator improves resolution."""
    split_header = split_header_with_separator(header, separator)
    if not split_header or split_header == safe_cell(header):
        return False
    if resolve_table_header_match_uncached(split_header, header_separator_chars=""):
        return True
    if SCAN_NLP_ENABLED:
        suggestion_fn = globals().get("nlp_header_suggestion")
        if callable(suggestion_fn):
            suggested_header, _confidence = suggestion_fn(split_header)
            if scan_nlp_header_suggestion_supported(split_header, suggested_header):
                return True
    return False


def infer_label_header_separator_chars(header):
    """Infer separators that are useful for resolving one header label."""
    header = safe_cell(header)
    if not header:
        return ""
    voted = []
    for separator in HEADER_WORD_SEPARATOR_CANDIDATES:
        if header_separator_split_supported(header, separator):
            voted.append(separator)
    return "".join(voted)


def infer_table_header_separator_chars(headers, sampled_rows=None, include_first_column_labels=False):
    """Infer word separators for this table's trusted header-like labels only."""
    labels = [safe_cell(header) for header in headers if safe_cell(header)]
    if include_first_column_labels and sampled_rows:
        labels.extend(
            safe_cell(row[0])
            for row in sampled_rows
            if row and safe_cell(row[0])
        )
    votes = Counter()
    for label in labels:
        # Do not learn separator semantics from generic placeholder headers.
        if rule_slug(label).startswith("column_"):
            continue
        for separator in infer_label_header_separator_chars(label):
            votes[separator] += 1
    return "".join(separator for separator in HEADER_WORD_SEPARATOR_CANDIDATES if votes.get(separator, 0) > 0)


def resolve_table_header_match(header, header_separator_chars=""):
    header = safe_cell(header)
    if not header:
        return None
    header_separator_chars = safe_cell(header_separator_chars)
    cache_key = (header, header_separator_chars)
    if cache_key in TABLE_HEADER_MATCH_CACHE:
        return TABLE_HEADER_MATCH_CACHE[cache_key]
    if re.fullmatch(r"column_\d+", header.casefold()):
        if len(TABLE_HEADER_MATCH_CACHE) >= 1024:
            TABLE_HEADER_MATCH_CACHE.clear()
        TABLE_HEADER_MATCH_CACHE[cache_key] = None
        return None
    auto_cache_key = None
    if not header_separator_chars:
        auto_cache_key = (header, "__auto__")
        if auto_cache_key in TABLE_HEADER_MATCH_CACHE:
            return TABLE_HEADER_MATCH_CACHE[auto_cache_key]
        header_separator_chars = infer_label_header_separator_chars(header)
    cache_key = (header, header_separator_chars)
    if cache_key in TABLE_HEADER_MATCH_CACHE:
        result = TABLE_HEADER_MATCH_CACHE[cache_key]
        if auto_cache_key is not None:
            TABLE_HEADER_MATCH_CACHE[auto_cache_key] = result
        return result

    result = None
    for leaf_candidate in structured_header_leaf_candidates(header):
        leaf_separator_chars = header_separator_chars or infer_label_header_separator_chars(leaf_candidate)
        leaf_result = resolve_table_header_match_uncached(leaf_candidate, header_separator_chars=leaf_separator_chars)
        if leaf_result:
            result = leaf_result
            break

    if result is None:
        result = resolve_table_header_match_uncached(header, header_separator_chars=header_separator_chars)

    if result is None:
        result = None
        if SCAN_NLP_ENABLED:
            suggestion_fn = globals().get("nlp_header_suggestion")
            if callable(suggestion_fn):
                suggested_header, _confidence = suggestion_fn(header)
                if scan_nlp_header_suggestion_supported(header, suggested_header):
                    metadata = {}
                    for _pattern, canonical_header, _order, candidate_metadata in TABLE_COLUMN_RULES.get("header_matchers", ()):
                        if canonical_header == suggested_header:
                            metadata = candidate_metadata
                            break
                    result = {
                        "canonical_header": suggested_header,
                        "metadata": metadata,
                    }

    if len(TABLE_HEADER_MATCH_CACHE) >= 1024:
        TABLE_HEADER_MATCH_CACHE.clear()
    TABLE_HEADER_MATCH_CACHE[cache_key] = result
    if auto_cache_key is not None:
        TABLE_HEADER_MATCH_CACHE[auto_cache_key] = result
    return result


def scan_nlp_header_suggestion_supported(header, suggested_header):
    if not suggested_header:
        return False
    if suggested_header in {"first name", "last name", "full name"} and is_non_person_name_header(header):
        return False

    label_text = semantic_label_text(header)
    label_tokens = semantic_label_tokens(header)

    if suggested_header == "DOB":
        if {"dob", "birth", "born"} & label_tokens or "date of birth" in label_text:
            return True
        return False

    if suggested_header == "person id":
        if rule_slug(header) in {"emplid", "empl_id"}:
            return True
        # Generic references may identify orders or documents rather than people.
        subject_tokens = {"candidate", "applicant", "person", "customer", "user", "employee", "staff", "student", "member"}
        id_tokens = {"id", "identifier", "number", "no", "ref", "reference"}
        return bool(subject_tokens & label_tokens and id_tokens & label_tokens)

    if suggested_header == "citizenship country":
        if any(phrase in label_text for phrase in (
            "country of nationality",
            "country of citizenship",
            "country of domicile",
        )):
            return True
        if {"nationality", "citizenship", "domicile", "citizen"} & label_tokens:
            return True
        return False

    if suggested_header in {"religion", "ethnicity", "disability", "marital status", "gender/sex"}:
        metadata = {}
        for _pattern, canonical_header, _order, candidate_metadata in TABLE_COLUMN_RULES.get("header_matchers", ()):
            if canonical_header == suggested_header:
                metadata = candidate_metadata
                break
        guard = metadata.get("context_guard")
        if guard and not header_context_supports_guard(header, guard):
            return False

    return True


def normalise_header(header, index, header_separator_chars=""):
    header = safe_cell(header)
    resolved = resolve_table_header(header, header_separator_chars=header_separator_chars)
    if resolved:
        return resolved
    return header if header else f"column_{index}"


HEADER_INFERENCE_SAMPLE_ROWS = 100
HEADER_INFERENCE_MIN_VALUES = 3
HEADER_INFERENCE_MIN_RATIO = 0.80
TABLE_SNIFF_DELIMITERS = ",\t;|~#"
HEADERLESS_PERSON_CONTEXT_HEADERS = {
    "email",
    "phone",
    "postcode",
    "person id",
    "student id",
    "staff id",
    "national insurance number",
    "nhs number",
    "passport",
    "payment card",
    "account number",
}
HEADERLESS_GENDER_VALUES = {
    "m", "male", "man", "f", "female", "woman", "non-binary", "non binary",
    "nb", "other", "prefer not to say", "unknown",
}


def full_cell_evidence_types(value):
    value = safe_cell(value)
    if not value:
        return set()

    evidence_types = set()
    for pattern in PATTERNS:
        if pattern.get("line_triggers") or pattern.get("labelled_fields_only"):
            continue
        for match in pattern["regex"].finditer(value):
            matched_text, start, end = extract_match_value(match)
            if value[:start].strip() or value[end:].strip():
                continue
            if validate_pattern_match(pattern, matched_text):
                evidence_types.add(evidence_type_for_match(pattern["evidence_type"], matched_text))
    if valid_uk_postcode(value):
        evidence_types.add("Postcode")
    return evidence_types


def dominant_column_evidence_type(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < HEADER_INFERENCE_MIN_VALUES:
        return ""

    counts = Counter()
    for value in values:
        counts.update(full_cell_evidence_types(value))
    if not counts:
        return ""

    ranked = counts.most_common()
    evidence_type, count = ranked[0]
    if evidence_type == "DOB":
        return ""
    if count / len(values) < HEADER_INFERENCE_MIN_RATIO:
        return ""
    if len(ranked) > 1 and ranked[1][1] == count:
        return ""
    return evidence_type


def inferred_header_for_evidence_type(evidence_type):
    if evidence_type in EMAIL_EVIDENCE_TYPES or evidence_type == "Email":
        return "email"
    return {
        "Phone": "phone",
        "DOB": "DOB",
        "Postcode": "postcode",
        "National Insurance Number": "national insurance number",
        "NHS Number": "nhs number",
        "Passport": "passport",
        "CAS Number": "CAS number",
        "HUSID": "HUSID",
        "SLC ID": "SLC ID",
        "Student SSN": "student SSN",
        "Payment Card": "payment card",
        "Bank Account": "account number",
    }.get(evidence_type, evidence_type)


def likely_full_name_column(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < HEADER_INFERENCE_MIN_VALUES:
        return False
    valid = [value for value in values if valid_name_value(value)]
    unique_ratio = len({normalize_name_value(value) for value in valid}) / len(values)
    return len(valid) / len(values) >= HEADER_INFERENCE_MIN_RATIO and unique_ratio >= 0.60


INFERRED_NAME_TOKEN_STOPWORDS = {
    "active", "inactive", "pending", "yes", "no", "true", "false", "male", "female", "other",
    "england", "scotland", "wales", "ireland", "christian", "muslim", "hindu", "sikh", "jewish",
}


def likely_name_token_column(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < HEADER_INFERENCE_MIN_VALUES:
        return False
    valid = [
        value
        for value in values
        if re.fullmatch(r"[A-Za-z][A-Za-z'\-]{1,39}", value)
        and value.casefold() not in INFERRED_NAME_TOKEN_STOPWORDS
    ]
    unique_ratio = len({value.casefold() for value in valid}) / len(values)
    return len(valid) / len(values) >= 0.90 and unique_ratio >= 0.60


def likely_gender_column(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < HEADER_INFERENCE_MIN_VALUES:
        return False
    normalized = [re.sub(r"\s+", " ", value.casefold()).strip() for value in values]
    matches = [value for value in normalized if value in HEADERLESS_GENDER_VALUES]
    distinct = set(matches)
    return len(matches) / len(values) >= 0.90 and 1 < len(distinct) <= 6


def likely_address_column(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < HEADER_INFERENCE_MIN_VALUES:
        return False
    matches = sum(has_searchable_address_component(value) for value in values)
    return matches / len(values) >= HEADER_INFERENCE_MIN_RATIO


WEAK_GENERIC_TABLE_HEADERS = {
    "", "id", "identifier", "number", "no", "ref", "reference", "code", "key", "value",
    "user", "member", "employee", "staff", "student", "person", "applicant", "record",
}


def is_weak_generic_table_header(header, header_separator_chars=""):
    slug = rule_slug(resolve_table_header(header, header_separator_chars=header_separator_chars) or header)
    return slug in {rule_slug(value) for value in WEAK_GENERIC_TABLE_HEADERS} or slug.startswith("column_")


def transposed_table_label_signal(label, header_separator_chars=""):
    """Return whether a first-column value looks like a field label.

    This is deliberately stricter than general header detection.  A transposed
    table should only be rotated when several first-column cells are actual
    known labels such as DOB, email, NHS number, nationality, etc.
    """
    label = safe_cell(label)
    if not label:
        return False
    resolved = resolve_table_header(label, header_separator_chars=header_separator_chars)
    if resolved and not is_weak_generic_table_header(resolved, header_separator_chars=header_separator_chars):
        return True
    return False


TRANSPOSED_HEADER_MARKERS = {
    "field", "fields", "attribute", "attributes", "label", "labels", "key", "keys",
    "question", "questions", "item", "items", "property", "properties", "name", "names",
}

TRANSPOSED_VALUE_HEADER_MARKERS = {
    "value", "values", "answer", "answers", "response", "responses", "data", "details",
    "person", "record", "subject", "student", "applicant", "candidate",
}

TRANSPOSED_RECORD_ID_HEADER_RE = re.compile(
    r"^(?:person|record|student|applicant|candidate|subject|member|user|contact|"
    r"customer|client|patient|p|r)[\s_\-#]*[A-Za-z0-9]+$",
    re.IGNORECASE,
)


def transposed_record_header_signal(value, header_separator_chars=""):
    """Return whether a column header looks like a record/person identifier.

    In a transposed table the first row normally looks like:

    ``Field, person_one, person_two``

    The cells after the first are record names/IDs, not data-field labels.  We
    also allow actual full names, because real transposed exports often use
    person names as the record headers.
    """
    value = safe_cell(value)
    if not value:
        return False
    if transposed_table_label_signal(value, header_separator_chars=header_separator_chars):
        return False
    slug = rule_slug(value)
    if slug in TRANSPOSED_VALUE_HEADER_MARKERS:
        return True
    if TRANSPOSED_RECORD_ID_HEADER_RE.fullmatch(value):
        return True
    if person_table_exact_match_count_for_value(value):
        return True
    if valid_name_value(value, require_full_name=True):
        return True
    evidence_types = full_cell_evidence_types(value)
    return has_id_evidence(evidence_types) or has_email_evidence(evidence_types)


def transposed_header_marker_signal(row, header_separator_chars=""):
    """Return whether a row looks like the marker/header row of a transposed table."""
    cells = [safe_cell(value) for value in row]
    if len(cells) < 2:
        return False
    first_slug = rule_slug(cells[0])
    if first_slug not in TRANSPOSED_HEADER_MARKERS:
        return False
    remaining = [value for value in cells[1:] if value]
    if not remaining:
        return False
    record_like = sum(1 for value in remaining if transposed_record_header_signal(value, header_separator_chars=header_separator_chars))
    return record_like >= 1 and record_like / len(remaining) >= 0.50


def split_leading_transposed_metadata(headers, sampled_rows, max_metadata_rows=5):
    """Skip report metadata rows before a transposed table header, when present."""
    candidate_rows = [list(headers)] + [list(row) for row in sampled_rows]
    for index, row in enumerate(candidate_rows[:max_metadata_rows]):
        following = candidate_rows[index + 1:]
        separator_chars = infer_table_header_separator_chars(row, following, include_first_column_labels=True)
        if not transposed_header_marker_signal(row, header_separator_chars=separator_chars):
            continue
        first_column_labels = [safe_cell(value[0]) for value in following if value and safe_cell(value[0])]
        signal_count = sum(
            1 for label in first_column_labels
            if transposed_table_label_signal(label, header_separator_chars=separator_chars)
        )
        if signal_count >= 2:
            return list(row), following, index
    return list(headers), [list(row) for row in sampled_rows], 0


def rotated_record_evidence_type_count(field_labels, record_values, header_separator_chars=""):
    """Count distinct person-evidence types visible in one rotated record."""
    evidence_types = set()
    for label, value in zip(field_labels, record_values):
        label = safe_cell(label)
        value = safe_cell(value)
        if not label or not value:
            continue
        resolved = resolve_table_header(label, header_separator_chars=header_separator_chars)
        if resolved and transposed_table_label_signal(label, header_separator_chars=header_separator_chars):
            metadata = resolve_table_header_metadata(label, header_separator_chars=header_separator_chars)
            for evidence_type in configured_evidence_types_from_header_metadata(metadata):
                if evidence_type:
                    evidence_types.add(evidence_type)
        evidence_types.update(full_cell_evidence_types(value))
    return len(evidence_types)


def table_orientation_assessment(headers, sampled_rows, header_separator_chars=""):
    """Assess whether a table is normal, transposed, headerless, or ambiguous.

    This is intentionally evidence-based.  A table is only rotated when several
    first-column cells resolve as known field labels and the columns after the
    first look like records/persons rather than normal data fields.  Ambiguous
    cases are left unrotated and downstream code treats them as review/raw
    fallback rather than trusted headers.
    """
    rows = [list(headers)] + [list(row) for row in sampled_rows]
    rows = [row for row in rows if any(safe_cell(value) for value in row)]
    if not rows:
        return {
            "orientation": "headerless",
            "confidence": "low",
            "reason": "empty_or_unusable_table_sample",
            "transpose": False,
            "review_only": True,
        }

    column_count = max(len(row) for row in rows)
    first_column_labels = [safe_cell(row[0]) for row in rows if row and safe_cell(row[0])]
    if len(rows) < 3 or len(first_column_labels) < 2:
        return {
            "orientation": "normal",
            "confidence": "low",
            "reason": f"insufficient_orientation_sample rows={len(rows)} first_column_labels={len(first_column_labels)}",
            "transpose": False,
            "review_only": False,
        }

    signal_count = sum(
        1 for label in first_column_labels
        if transposed_table_label_signal(label, header_separator_chars=header_separator_chars)
    )
    signal_ratio = signal_count / len(first_column_labels) if first_column_labels else 0.0

    record_headers = [safe_cell(value) for value in rows[0][1:]]
    non_empty_record_headers = [value for value in record_headers if value]
    record_header_signals = sum(
        1 for value in non_empty_record_headers
        if transposed_record_header_signal(value, header_separator_chars=header_separator_chars)
    )
    record_header_ratio = (
        record_header_signals / len(non_empty_record_headers)
        if non_empty_record_headers else 0.0
    )
    field_like_record_headers = sum(
        1 for value in non_empty_record_headers
        if transposed_table_label_signal(value, header_separator_chars=header_separator_chars)
    )
    field_like_record_header_ratio = (
        field_like_record_headers / len(non_empty_record_headers)
        if non_empty_record_headers else 0.0
    )

    field_label_rows = [
        row for row in rows
        if row and safe_cell(row[0]) and transposed_table_label_signal(row[0], header_separator_chars=header_separator_chars)
    ]
    field_labels = [safe_cell(row[0]) for row in field_label_rows]
    record_count = max(column_count - 1, 0)
    rotated_evidence_counts = []
    for record_index in range(record_count):
        record_values = [
            safe_cell(row[record_index + 1]) if record_index + 1 < len(row) else ""
            for row in field_label_rows
        ]
        rotated_evidence_counts.append(
            rotated_record_evidence_type_count(field_labels, record_values, header_separator_chars=header_separator_chars)
        )
    strong_rotated_records = sum(1 for count in rotated_evidence_counts if count >= 2)
    any_rotated_record_signal = max(rotated_evidence_counts or [0])

    reasons = (
        f"first_column_label_signals={signal_count}/{len(first_column_labels)} "
        f"record_header_signals={record_header_signals}/{len(non_empty_record_headers)} "
        f"field_like_record_headers={field_like_record_headers}/{len(non_empty_record_headers)} "
        f"rotated_records_with_2plus_evidence={strong_rotated_records}/{record_count}"
    )

    if column_count == 2:
        second_header = non_empty_record_headers[0] if non_empty_record_headers else ""
        value_header = (
            rule_slug(second_header) in TRANSPOSED_VALUE_HEADER_MARKERS
            or transposed_record_header_signal(second_header, header_separator_chars=header_separator_chars)
        )
        if signal_count >= 2 and signal_ratio >= 0.50 and value_header and any_rotated_record_signal >= 2:
            return {
                "orientation": "transposed",
                "confidence": "medium",
                "reason": "field_value_form; " + reasons,
                "transpose": True,
                "review_only": False,
            }
        if signal_count >= 2 and signal_ratio >= 0.50:
            return {
                "orientation": "ambiguous",
                "confidence": "medium",
                "reason": "field_value_like_but_record_value_header_unclear; " + reasons,
                "transpose": False,
                "review_only": True,
            }

    if column_count >= 3 and signal_count >= 2 and signal_ratio >= 0.35:
        enough_record_headers = record_header_ratio >= 0.50
        enough_rotated_records = strong_rotated_records >= max(1, min(record_count, 2))
        record_headers_are_not_fields = field_like_record_header_ratio <= 0.25
        if enough_record_headers and enough_rotated_records and record_headers_are_not_fields:
            confidence = "high" if signal_ratio >= 0.50 and record_header_ratio >= 0.75 else "medium"
            return {
                "orientation": "transposed",
                "confidence": confidence,
                "reason": reasons,
                "transpose": True,
                "review_only": False,
            }
        if field_like_record_header_ratio >= 0.25 or record_header_ratio >= 0.25 or strong_rotated_records:
            return {
                "orientation": "ambiguous",
                "confidence": "medium",
                "reason": "partial_transposed_signal; " + reasons,
                "transpose": False,
                "review_only": True,
            }

    initial_resolved_headers = [normalise_header(h, i + 1, header_separator_chars=header_separator_chars) for i, h in enumerate(headers)]
    if looks_like_table_headers(initial_resolved_headers, header_separator_chars=header_separator_chars):
        return {
            "orientation": "normal",
            "confidence": "high",
            "reason": "source_header_row_matches_known_table_headers",
            "transpose": False,
            "review_only": False,
        }
    return {
        "orientation": "headerless",
        "confidence": "low",
        "reason": "no_reliable_header_or_transposed_signal; " + reasons,
        "transpose": False,
        "review_only": False,
    }


def looks_like_transposed_table(headers, sampled_rows):
    """Detect tables where rows are fields and columns are records."""
    header_separator_chars = infer_table_header_separator_chars(headers, sampled_rows, include_first_column_labels=True)
    return table_orientation_assessment(headers, sampled_rows, header_separator_chars=header_separator_chars)["orientation"] == "transposed"


def transpose_field_labelled_table(headers, sampled_rows, remaining_rows):
    """Rotate a field-labelled/transposed table into conventional rows."""
    source_rows = [list(headers)] + [list(row) for row in sampled_rows] + [list(row) for row in remaining_rows]
    source_rows = [row for row in source_rows if any(safe_cell(value) for value in row)]
    field_headers = [safe_cell(row[0]) or f"field_{index + 1}" for index, row in enumerate(source_rows)]
    record_count = max([len(row) for row in source_rows] or [1]) - 1
    transposed_rows = []
    for record_index in range(record_count):
        row = [
            safe_cell(source_row[record_index + 1]) if record_index + 1 < len(source_row) else ""
            for source_row in source_rows
        ]
        if any(row):
            transposed_rows.append(row)
    return field_headers, iter(transposed_rows)


def infer_table_headers(sample_rows, base_headers=None, trust_base_headers=False, header_separator_chars=""):
    """Infer column roles from table values, using person-table matches first.

    This runs on every table, not only headerless tables.  The precedence is:
    1. exact value matches from --person-tables;
    2. explicit/specific source headers, when present;
    3. regex/value-based column inference;
    4. generic column_N fallback.

    Generic headers such as ``ID`` are treated as weak.  They are only resolved
    to a specific identifier class when the values match a specific
    --person-tables column/type.
    """
    rows = [list(row) for row in sample_rows if any(safe_cell(value) for value in row)]
    if len(rows) < HEADER_INFERENCE_MIN_VALUES:
        if trust_base_headers and base_headers:
            return [safe_cell(header) or f"column_{index + 1}" for index, header in enumerate(base_headers)]
        return []

    column_count = max(
        [len(row) for row in rows] + ([len(base_headers)] if base_headers else [0])
    )
    columns = [
        [row[index] if index < len(row) else "" for row in rows]
        for index in range(column_count)
    ]
    headers = [f"column_{index + 1}" for index in range(column_count)]
    inferred_indexes = set()

    # 1. Strongest signal: the values themselves match a known person-table
    # column/type.  This is deliberately applied to all tables, including ones
    # with headers, because local identifier values are more informative than
    # weak labels such as ID, number, user, or reference.
    for index, values in enumerate(columns):
        source_header = safe_cell(base_headers[index]) if trust_base_headers and base_headers and index < len(base_headers) else ""
        if trust_base_headers and source_header and not is_weak_generic_table_header(source_header, header_separator_chars=header_separator_chars):
            continue
        inferred_header = infer_header_from_person_table_values(values)
        if inferred_header:
            headers[index] = inferred_header
            inferred_indexes.add(index)

    # 2. If value matching could not type a column, keep explicit source
    # headers where they are useful.  Weak generic labels are kept only as a
    # last resort, so ``ID`` does not accidentally become a specific staff or
    # student identifier unless the values support that interpretation.
    if trust_base_headers and base_headers:
        for index in range(column_count):
            if index in inferred_indexes:
                continue
            source_header = safe_cell(base_headers[index]) if index < len(base_headers) else ""
            if source_header and not is_weak_generic_table_header(source_header, header_separator_chars=header_separator_chars):
                headers[index] = source_header
                inferred_indexes.add(index)

    # 3. Existing regex/value-based column inference for unresolved columns.
    for index, values in enumerate(columns):
        if index in inferred_indexes:
            continue
        evidence_type = dominant_column_evidence_type(values)
        if evidence_type:
            headers[index] = inferred_header_for_evidence_type(evidence_type)
            inferred_indexes.add(index)

    # Headerless tables can still be person records when value columns form a
    # coherent row-level pattern.  Only infer names/demographics when another
    # already-inferred column provides person context, e.g. email + first +
    # last + gender.  Do not infer DOB here: date-shaped columns are too often
    # course, event, or transaction dates without an explicit birth label.
    if not trust_base_headers:
        inferred_header_set = {
            normalise_header(headers[index], index + 1, header_separator_chars=header_separator_chars)
            for index in inferred_indexes
        }
        has_person_context = bool(inferred_header_set & HEADERLESS_PERSON_CONTEXT_HEADERS)
        if has_person_context:
            for index, values in enumerate(columns):
                if index in inferred_indexes:
                    continue
                if likely_gender_column(values):
                    headers[index] = "gender/sex"
                    inferred_indexes.add(index)

            for index, values in enumerate(columns):
                if index in inferred_indexes:
                    continue
                if likely_full_name_column(values):
                    headers[index] = "full name"
                    inferred_indexes.add(index)

            for index in range(column_count - 1):
                if index in inferred_indexes or index + 1 in inferred_indexes:
                    continue
                if likely_name_token_column(columns[index]) and likely_name_token_column(columns[index + 1]):
                    headers[index] = "first name"
                    headers[index + 1] = "last name"
                    inferred_indexes.add(index)
                    inferred_indexes.add(index + 1)

    # Outside the contextual headerless branch above, do not infer person-name
    # columns in unknown/weakly-labelled tables from generic "looks like a name"
    # heuristics. Unknown name columns are otherwise only typed when supported
    # by several known person-table values, via step 1 above.

    postcode_indexes = {
        index for index, header in enumerate(headers)
        if normalise_header(header, index + 1, header_separator_chars=header_separator_chars) == "postcode"
    }
    if postcode_indexes:
        for index, values in enumerate(columns):
            if index in inferred_indexes:
                continue
            if likely_address_column(values) and min(abs(index - postcode) for postcode in postcode_indexes) <= 3:
                headers[index] = "address1"
                inferred_indexes.add(index)

    # 4. If a table has trusted headers, retain weak/generic labels for any
    # columns that remain unresolved.  Otherwise require at least one inferred
    # column before converting the table to labelled evidence.
    if trust_base_headers and base_headers:
        for index in range(column_count):
            if index in inferred_indexes:
                continue
            source_header = safe_cell(base_headers[index]) if index < len(base_headers) else ""
            if source_header:
                headers[index] = source_header
        return headers

    if not inferred_indexes:
        return []
    return headers


def table_rows_to_labelled_text(headers, rows, table_name="table", max_rows=0):
    """Convert table rows into label:value lines so regex context can use headings.

    This is intentionally plain text. It means the existing regex, finding and
    candidate logic can treat a table row as a person-level record, but each cell
    now carries its column heading as context, e.g. `DOB: 1999-04-12` rather than
    a bare date value.

    Column typing now runs value-first on every table:
    - exact matches against --person-tables values;
    - explicit/specific headers for unresolved columns;
    - regex/value-based inference for anything still unresolved.
    """
    rows = iter(rows)
    sampled_data_rows = list(islice(rows, HEADER_INFERENCE_SAMPLE_ROWS))
    original_headers = list(headers)

    original_headers, sampled_data_rows, _metadata_rows_skipped = split_leading_transposed_metadata(
        original_headers,
        sampled_data_rows,
    )
    header_separator_chars = infer_table_header_separator_chars(
        original_headers,
        sampled_data_rows,
        include_first_column_labels=True,
    )
    orientation = table_orientation_assessment(
        original_headers,
        sampled_data_rows,
        header_separator_chars=header_separator_chars,
    )
    force_review_headerless = orientation["orientation"] == "ambiguous"

    if force_review_headerless:
        return ""

    if orientation["transpose"]:
        original_headers, rows = transpose_field_labelled_table(original_headers, sampled_data_rows, rows)
        sampled_data_rows = list(islice(rows, HEADER_INFERENCE_SAMPLE_ROWS))

    initial_resolved_headers = [
        normalise_header(h, i + 1, header_separator_chars=header_separator_chars)
        for i, h in enumerate(original_headers)
    ]
    recognized_headers = looks_like_table_headers(initial_resolved_headers, header_separator_chars=header_separator_chars)
    headers_contain_evidence_values = any(full_cell_evidence_types(value) for value in original_headers)
    headers_contain_person_table_values = any(person_table_exact_match_count_for_value(value) for value in original_headers)

    # If the first row looks like actual evidence/person data, treat it as a
    # headerless data row even if the CSV reader handed it to us as headers.
    # Otherwise, keep it as the source header row.
    treat_first_row_as_data = (
        force_review_headerless
        or (not recognized_headers)
        or headers_contain_evidence_values
        or headers_contain_person_table_values
    )
    if treat_first_row_as_data:
        sampled_rows = [original_headers] + sampled_data_rows
        base_headers = [f"column_{index + 1}" for index in range(max([len(row) for row in sampled_rows] or [len(original_headers)]))]
        trust_base_headers = False
        rows = chain(sampled_rows, rows)
    else:
        sampled_rows = sampled_data_rows
        base_headers = original_headers
        trust_base_headers = True
        rows = chain(sampled_data_rows, rows)

    table_header_separator_chars = header_separator_chars if trust_base_headers or orientation["transpose"] else ""
    inferred_headers = infer_table_headers(
        sampled_rows,
        base_headers=base_headers,
        trust_base_headers=trust_base_headers,
        header_separator_chars=table_header_separator_chars,
    )
    if not inferred_headers:
        return ""

    original_headers = inferred_headers
    resolved_headers = [
        normalise_header(h, i + 1, header_separator_chars=table_header_separator_chars)
        for i, h in enumerate(original_headers)
    ]
    display_headers = [
        safe_cell(original_headers[index]) or resolved_headers[index] or f"column_{index + 1}"
        for index in range(len(original_headers))
    ]

    header_metadata = [
        resolve_table_header_metadata(h, header_separator_chars=table_header_separator_chars)
        for h in original_headers
    ]

    parts = [f"\n--- TABLE: {table_name} ---\n"]
    emitted = 0
    for row_index, row in enumerate(rows, start=1):
        if max_rows and emitted >= max_rows:
            break
        cells = [safe_cell(v) for v in row]
        if not any(cells):
            continue
        labelled = []
        for col_index, header in enumerate(display_headers):
            value = cells[col_index] if col_index < len(cells) else ""
            if value:
                labelled.append(f"{header}: {value}")
        name_fields = [
            (original_headers[index], cells[index] if index < len(cells) else "", header_metadata[index])
            for index in range(len(original_headers))
        ]
        for name in assembled_names_from_fields(name_fields):
            labelled.append(f"full name: {name['matched_text']}")
        if labelled:
            parts.append(f"TABLE_ROW {table_name} #{row_index} | " + " | ".join(labelled))
            emitted += 1
    return "\n".join(parts)


EMBEDDED_TABLE_DELIMITERS = (",", "\t", ";", "|", "~", "#", "_", "-")
SAFE_EMBEDDED_TABLE_DELIMITERS = (",", "\t", ";", "|", "~", "#")
STRICT_EMBEDDED_TABLE_DELIMITERS = {"_", "-"}
EMBEDDED_TABLE_VALIDATION_SAMPLE_ROWS = HEADER_INFERENCE_SAMPLE_ROWS


def detected_delimiter_for_line(line):
    """Return candidate delimiters that could split one plain-text line."""
    line = "" if line is None else str(line).strip("\r\n")
    if not line:
        return {}

    # Strong separators are evaluated first.  If a line contains a proper
    # table separator such as ``~`` or ``|``, do not also treat hyphens or
    # underscores inside values/IDs as competing table delimiters.
    candidates = {}
    for delimiter in SAFE_EMBEDDED_TABLE_DELIMITERS:
        count = line.count(delimiter)
        # Embedded table detection is intentionally stricter than whole-file
        # CSV sniffing.  Requiring at least three fields avoids treating casual
        # punctuation in prose as a table row.
        if count >= 2:
            candidates[delimiter] = count
    if candidates:
        return candidates

    for delimiter in STRICT_EMBEDDED_TABLE_DELIMITERS:
        count = line.count(delimiter)
        if count >= 2:
            candidates[delimiter] = count
    return candidates


def quick_split_embedded_rows(lines, delimiter):
    return [
        [safe_cell(part) for part in line.split(delimiter)]
        for line in lines
    ]


def embedded_name_token_column(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < 2:
        return False
    matches = [
        value
        for value in values
        if re.fullmatch(r"[A-Za-z][A-Za-z'\-]{1,39}", value)
        and value.casefold() not in INFERRED_NAME_TOKEN_STOPWORDS
    ]
    return len(matches) / len(values) >= 0.90


def embedded_gender_column(values):
    values = [safe_cell(value) for value in values if safe_cell(value)]
    if len(values) < 2:
        return False
    normalized = [re.sub(r"\s+", " ", value.casefold()).strip() for value in values]
    matches = [value for value in normalized if value in HEADERLESS_GENDER_VALUES]
    return len(matches) / len(values) >= 0.90


def embedded_delimited_block_is_table(block_lines, delimiter, field_count):
    sample_lines = list(block_lines[:EMBEDDED_TABLE_VALIDATION_SAMPLE_ROWS])
    rows = quick_split_embedded_rows(sample_lines, delimiter)
    if len(rows) < 2 or not rows:
        return False
    if any(len(row) != field_count for row in rows):
        return False
    cells = [cell for row in rows for cell in row]
    non_empty_ratio = len([cell for cell in cells if cell]) / max(1, len(cells))
    if non_empty_ratio < 0.60:
        return False
    if field_count >= 3:
        columns = [
            [row[index] if index < len(row) else "" for row in rows]
            for index in range(field_count)
        ]
        has_gender = any(embedded_gender_column(column) for column in columns)
        has_name_pair = any(
            embedded_name_token_column(columns[index]) and embedded_name_token_column(columns[index + 1])
            for index in range(field_count - 1)
        )
        if has_gender and has_name_pair:
            return True
    if delimiter in STRICT_EMBEDDED_TABLE_DELIMITERS:
        return False

    if len(rows) >= 3 and field_count >= 3:
        return True

    inferred = infer_table_headers(
        rows,
        base_headers=[f"column_{index + 1}" for index in range(field_count)],
        trust_base_headers=False,
    )
    return any(
        safe_cell(header) and not safe_cell(header).startswith("column_")
        for header in inferred
    )


def detect_embedded_delimited_blocks(text, min_rows=2):
    """Find repeated delimiter-shaped table blocks inside otherwise plain text."""
    lines = text.splitlines()
    blocks = []
    index = 0
    while index < len(lines):
        best = None
        line_candidates = detected_delimiter_for_line(lines[index])
        for delimiter, delimiter_count in line_candidates.items():
            end = index + 1
            while end < len(lines):
                next_candidates = detected_delimiter_for_line(lines[end])
                if next_candidates.get(delimiter) != delimiter_count:
                    break
                end += 1
            row_count = end - index
            if row_count >= min_rows:
                sample_end = min(end, index + EMBEDDED_TABLE_VALIDATION_SAMPLE_ROWS)
                block_sample = lines[index:sample_end]
                if not embedded_delimited_block_is_table(block_sample, delimiter, delimiter_count + 1):
                    continue
                score = (row_count, delimiter_count)
                if best is None or score > best[0]:
                    best = (score, delimiter, delimiter_count, end)
        if best:
            _score, delimiter, delimiter_count, end = best
            blocks.append({
                "start": index,
                "end": end,
                "delimiter": delimiter,
                "field_count": delimiter_count + 1,
            })
            index = end
        else:
            index += 1
    return blocks


def generic_headerless_table_text(rows, table_name, max_rows=0):
    """Emit generic column_N labelled rows for a structurally detected table."""
    rows = [list(row) for row in rows if any(safe_cell(value) for value in row)]
    if not rows:
        return ""
    column_count = max(len(row) for row in rows)
    headers = [f"column_{index + 1}" for index in range(column_count)]
    parts = [f"\n--- TABLE: {table_name} ---\n"]
    emitted = 0
    for row_index, row in enumerate(rows, start=1):
        if max_rows and emitted >= max_rows:
            break
        cells = [safe_cell(value) for value in row]
        labelled = [
            f"{headers[index]}: {cells[index]}"
            for index in range(min(column_count, len(cells)))
            if cells[index]
        ]
        if labelled:
            parts.append(f"TABLE_ROW {table_name} #{row_index} | " + " | ".join(labelled))
            emitted += 1
    return "\n".join(parts)


def rows_to_labelled_text_with_headers(headers, rows, table_name, max_rows=0, header_separator_chars=""):
    """Emit labelled table rows using already-resolved/inferred headers.

    This is the cheap output path for embedded headerless tables.  It avoids
    rerunning full orientation/header detection for every small block while
    preserving name assembly and column labels.
    """
    rows = [list(row) for row in rows if any(safe_cell(value) for value in row)]
    headers = [safe_cell(header) for header in headers]
    if not rows or not headers:
        return ""

    resolved_headers = [
        normalise_header(header, index + 1, header_separator_chars=header_separator_chars)
        for index, header in enumerate(headers)
    ]
    display_headers = [
        safe_cell(headers[index]) or resolved_headers[index] or f"column_{index + 1}"
        for index in range(len(headers))
    ]
    header_metadata = [
        resolve_table_header_metadata(header, header_separator_chars=header_separator_chars)
        for header in headers
    ]

    parts = [f"\n--- TABLE: {table_name} ---\n"]
    emitted = 0
    for row_index, row in enumerate(rows, start=1):
        if max_rows and emitted >= max_rows:
            break
        cells = [safe_cell(value) for value in row]
        labelled = []
        for col_index, header in enumerate(display_headers):
            value = cells[col_index] if col_index < len(cells) else ""
            if value:
                labelled.append(f"{header}: {value}")
        name_fields = [
            (headers[index], cells[index] if index < len(cells) else "", header_metadata[index])
            for index in range(len(headers))
        ]
        for name in assembled_names_from_fields(name_fields):
            labelled.append(f"full name: {name['matched_text']}")
        if labelled:
            parts.append(f"TABLE_ROW {table_name} #{row_index} | " + " | ".join(labelled))
            emitted += 1
    return "\n".join(parts)


def embedded_first_row_looks_like_header(first_row, data_rows):
    """Return whether the first row of an embedded block is a real header row."""
    first_row = [safe_cell(value) for value in first_row]
    if not first_row or not any(first_row):
        return False
    if any(full_cell_evidence_types(value) for value in first_row if value):
        return False
    if any(person_table_exact_match_count_for_value(value) for value in first_row if value):
        return False
    header_separator_chars = infer_table_header_separator_chars(
        first_row,
        data_rows[:HEADER_INFERENCE_SAMPLE_ROWS],
        include_first_column_labels=False,
    )
    resolved_headers = [
        normalise_header(header, index + 1, header_separator_chars=header_separator_chars)
        for index, header in enumerate(first_row)
    ]
    return looks_like_table_headers(resolved_headers, header_separator_chars=header_separator_chars)


def embedded_headerless_table_text(rows, table_name, schema_cache=None, max_rows=0):
    """Label an embedded headerless table block with cheap per-file schema reuse."""
    rows = [list(row) for row in rows if any(safe_cell(value) for value in row)]
    if not rows:
        return ""
    field_count = max(len(row) for row in rows)
    cache_key = field_count
    headers = schema_cache.get(cache_key) if schema_cache is not None else None
    if headers is None:
        sampled_rows = rows[:HEADER_INFERENCE_SAMPLE_ROWS]
        headers = infer_table_headers(
            sampled_rows,
            base_headers=[f"column_{index + 1}" for index in range(field_count)],
            trust_base_headers=False,
        )
        if schema_cache is not None:
            schema_cache[cache_key] = headers
    if headers:
        return rows_to_labelled_text_with_headers(headers, rows, table_name, max_rows=max_rows)
    return generic_headerless_table_text(rows, table_name, max_rows=max_rows)


def parse_delimited_lines(lines, delimiter):
    return list(csv.reader(io.StringIO("\n".join(lines), newline=""), delimiter=delimiter))


def labelled_embedded_delimited_blocks(text, table_name):
    """Replace repeated delimiter-shaped plain-text blocks with labelled rows."""
    lines = text.splitlines()
    blocks = detect_embedded_delimited_blocks(text)
    if not blocks:
        return ""

    output = []
    cursor = 0
    headerless_schema_cache = {}
    for block_index, block in enumerate(blocks, start=1):
        output.extend(lines[cursor:block["start"]])
        block_lines = lines[block["start"]:block["end"]]
        parsed_rows = parse_delimited_lines(block_lines, block["delimiter"])
        if parsed_rows:
            block_table_name = f"{table_name} embedded_table_{block_index}"
            if embedded_first_row_looks_like_header(parsed_rows[0], parsed_rows[1:]):
                headers = parsed_rows[0]
                data_rows = parsed_rows[1:]
                labelled = table_rows_to_labelled_text(
                    headers,
                    iter(data_rows),
                    block_table_name,
                )
            else:
                labelled = embedded_headerless_table_text(
                    parsed_rows,
                    block_table_name,
                    schema_cache=headerless_schema_cache,
                )
            if not labelled:
                labelled = generic_headerless_table_text(
                    parsed_rows,
                    block_table_name,
                )
            if labelled:
                output.append(labelled)
            else:
                output.extend(block_lines)
        cursor = block["end"]
    output.extend(lines[cursor:])
    return "\n".join(output)



STRUCTURED_PERSON_NODE_RE = re.compile(
    r"\b(?:person|people|student|staff|employee|worker|patient|participant|subject|applicant|candidate|"
    r"dependa?nt|passenger|member|user|customer|client|reviewer|manager|contact|guardian|parent|child|"
    r"record|row|entry)\b",
    re.IGNORECASE,
)

STRUCTURED_FRAGMENT_NODE_RE = re.compile(
    r"\b(?:address|addresses|contact[_\s-]?details|contactdetails|contacts?|phones?|emails?|names?|"
    r"documents?|identity|identifiers?|demographics?|passport|visa|immigration|payment|bank|account|"
    r"metadata|audit|history|logs?)\b",
    re.IGNORECASE,
)

STRUCTURED_STRONG_FIELD_RE = re.compile(
    r"\b(?:e[-_\s]?mail|email|mail|student[_\s-]?(?:id|number|no)|staff[_\s-]?(?:id|number|no)|"
    r"person[_\s-]?(?:id|number|no)|employee[_\s-]?(?:id|number|no)|assignment[_\s-]?(?:id|number|no)|"
    r"user(?:name|[_\s-]?id)?|passport|national[_\s-]?(?:id|identifier|number|no)|nhs|ucas|cas|"
    r"visa|brp|id[_\s-]?(?:number|no|value)?|identifier)\b",
    re.IGNORECASE,
)

STRUCTURED_SUPPORT_FIELD_RE = re.compile(
    r"\b(?:full[_\s-]?name|forename|first[_\s-]?name|given[_\s-]?name|surname|last[_\s-]?name|family[_\s-]?name|"
    r"name|dob|date[_\s-]?of[_\s-]?birth|birth[_\s-]?date|phone|telephone|mobile|address|postcode|"
    r"post[_\s-]?code|postal[_\s-]?code|sex|gender|nationality)\b",
    re.IGNORECASE,
)

STRUCTURED_SAFE_CONTEXT_RE = re.compile(
    r"\b(?:application|case|record|source|system|export|batch|file|faculty|school|department|status|"
    r"created|updated|timestamp|date|organisation|organization|role|type|category|reference|ref|"
    r"job|position|pay[_\s-]?grade|pool|workflow|form|submission|version)\b",
    re.IGNORECASE,
)

STRUCTURED_UNSAFE_CONTEXT_RE = re.compile(
    r"\b(?:e[-_\s]?mail|email|student|staff|person|employee|assignment|username|passport|national|nhs|"
    r"ucas|cas|visa|brp|identifier|full[_\s-]?name|forename|first[_\s-]?name|surname|last[_\s-]?name|"
    r"name|dob|birth|phone|telephone|mobile|address|postcode|postal|sex|gender|nationality)\b",
    re.IGNORECASE,
)

STRUCTURED_EMAIL_VALUE_RE = re.compile(r"\b[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}\b", re.IGNORECASE)


# =============================================================================

# Merged from extraction(3).py

# =============================================================================

"""Structured data parsing and file/text extraction."""



def sync_runtime_state():
    return None


def structured_humanise_token(value):
    value = str(value).lstrip("@").replace("_", " ").replace("-", " ")
    value = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", value)
    return value


def structured_path_label(path_parts):
    labels = []
    for part in path_parts:
        part = str(part)
        if part.startswith("[") and part.endswith("]"):
            continue
        if part in {"$", "#text"}:
            continue
        if part.startswith("@"):
            part = part[1:]
        labels.append(part)
    return ".".join(labels)


def structured_path_text(path_parts):
    if not path_parts:
        return "$"
    out = ""
    for index, part in enumerate(path_parts):
        part = str(part)
        if index == 0:
            out = part
        elif part.startswith("["):
            out += part
        else:
            out += "." + part
    return out or "$"


def structured_is_scalar(value):
    return value is None or isinstance(value, (str, int, float, bool))


def structured_scalar_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return safe_cell(value)


def structured_node_key(path_parts):
    for part in reversed(path_parts):
        part = str(part)
        if part.startswith("[") or part in {"$", "#text"}:
            continue
        return part.lstrip("@")
    return ""


def structured_key_blob(path_parts):
    return " ".join(structured_humanise_token(part) for part in path_parts)


def structured_field_score(label, value):
    blob = structured_humanise_token(safe_cell(label).replace(".", " "))
    value = safe_cell(value)
    strong = 0
    supporting = 0

    if STRUCTURED_STRONG_FIELD_RE.search(blob):
        strong += 1
    if STRUCTURED_SUPPORT_FIELD_RE.search(blob):
        supporting += 1
    if STRUCTURED_EMAIL_VALUE_RE.search(value):
        strong += 1
    if UK_POSTCODE_RE.search(value):
        supporting += 1
    if re.search(r"\b(?:dob|birth|date)\b", blob, re.IGNORECASE) and valid_contextual_date(value):
        supporting += 1

    return strong, supporting


def structured_direct_scalar_fields(node, path_parts):
    fields = []
    if isinstance(node, dict):
        for key, value in node.items():
            child_path = path_parts + [str(key)]
            if structured_is_scalar(value):
                text = structured_scalar_text(value)
                if text:
                    fields.append((child_path, structured_path_label(child_path), text))
    return fields


def structured_all_scalar_fields(node, path_parts, max_depth, current_depth=0):
    fields = []
    if current_depth > max_depth:
        return fields

    if structured_is_scalar(node):
        text = structured_scalar_text(node)
        if text:
            fields.append((path_parts, structured_path_label(path_parts), text))
        return fields

    if isinstance(node, dict):
        for key, value in node.items():
            fields.extend(structured_all_scalar_fields(value, path_parts + [str(key)], max_depth, current_depth + 1))
    elif isinstance(node, list):
        for index, value in enumerate(node):
            fields.extend(structured_all_scalar_fields(value, path_parts + [f"[{index}]"], max_depth, current_depth + 1))

    return fields


def structured_evidence_score(fields):
    strong = 0
    supporting = 0
    for _path, label, value in fields:
        field_strong, field_supporting = structured_field_score(label, value)
        strong += field_strong
        supporting += field_supporting
    return strong, supporting, strong * 3 + supporting


def structured_node_role(path_parts, node, max_depth):
    key = structured_node_key(path_parts)
    key_blob = structured_humanise_token(key)

    direct_fields = structured_direct_scalar_fields(node, path_parts)
    descendant_fields = structured_all_scalar_fields(node, path_parts, max_depth)
    direct_strong, direct_supporting, direct_score = structured_evidence_score(direct_fields)
    descendant_strong, descendant_supporting, descendant_score = structured_evidence_score(descendant_fields)

    semantic_person = bool(STRUCTURED_PERSON_NODE_RE.search(key_blob))
    semantic_fragment = bool(STRUCTURED_FRAGMENT_NODE_RE.search(key_blob))

    has_direct_name_like = any(
        re.search(r"\b(?:name|forename|surname|first[_\s-]?name|last[_\s-]?name)\b", label, re.IGNORECASE)
        for _path, label, _value in direct_fields
    )
    has_direct_identity = direct_strong > 0 or (has_direct_name_like and direct_supporting > 0)

    return {
        "key": key,
        "semantic_person": semantic_person,
        "semantic_fragment": semantic_fragment,
        "direct_fields": direct_fields,
        "descendant_fields": descendant_fields,
        "direct_strong": direct_strong,
        "direct_supporting": direct_supporting,
        "direct_score": direct_score,
        "descendant_strong": descendant_strong,
        "descendant_supporting": descendant_supporting,
        "descendant_score": descendant_score,
        "has_direct_identity": has_direct_identity,
    }


def structured_is_collection_item(path_parts):
    return bool(path_parts and str(path_parts[-1]).startswith("[") and str(path_parts[-1]).endswith("]"))


def structured_is_record_candidate(node, path_parts, max_depth):
    if structured_is_scalar(node):
        return False

    role = structured_node_role(path_parts, node, max_depth)
    if role["descendant_score"] < 2:
        return False

    in_collection = structured_is_collection_item(path_parts)

    # Fragment nodes such as address/contactDetails/identifiers are usually evidence
    # substructures, not record boundaries. They are folded into the nearest record.
    if role["semantic_fragment"] and not in_collection:
        return False

    if in_collection and role["descendant_score"] >= 2:
        return True

    if role["semantic_person"] and role["descendant_score"] >= 2:
        return True

    if role["has_direct_identity"] and role["direct_score"] >= 2:
        return True

    if role["direct_score"] >= 3:
        return True

    # Fallback for badly named but clearly entity-bearing objects.
    if role["descendant_strong"] >= 1 and role["descendant_supporting"] >= 2 and not role["semantic_fragment"]:
        return True

    return False


def structured_child_containers(node, path_parts):
    children = []
    if isinstance(node, dict):
        for key, value in node.items():
            if isinstance(value, (dict, list)):
                children.append((path_parts + [str(key)], value))
    elif isinstance(node, list):
        for index, value in enumerate(node):
            if isinstance(value, (dict, list)):
                children.append((path_parts + [f"[{index}]"], value))
    return children


def structured_nearest_record_paths(node, path_parts, max_depth, depth=0):
    record_paths = set()
    if depth > max_depth or structured_is_scalar(node):
        return record_paths

    if depth > 0 and structured_is_record_candidate(node, path_parts, max_depth):
        record_paths.add(tuple(path_parts))
        return record_paths

    for child_path, child in structured_child_containers(node, path_parts):
        record_paths.update(structured_nearest_record_paths(child, child_path, max_depth, depth + 1))

    return record_paths


def structured_record_child_paths(node, path_parts, max_depth):
    record_paths = set()
    for child_path, child in structured_child_containers(node, path_parts):
        record_paths.update(structured_nearest_record_paths(child, child_path, max_depth, 1))
    return record_paths


def structured_safe_context_fields(node, path_parts):
    context = []
    for child_path, label, value in structured_direct_scalar_fields(node, path_parts):
        blob = structured_key_blob(child_path)
        if STRUCTURED_SAFE_CONTEXT_RE.search(blob) and not STRUCTURED_UNSAFE_CONTEXT_RE.search(blob):
            context.append((child_path, label, value))
    return context


def structured_should_skip_subtree(child_path, record_child_paths):
    child_tuple = tuple(child_path)
    for record_path in record_child_paths:
        if len(child_tuple) <= len(record_path) and record_path[: len(child_tuple)] == child_tuple:
            return True
    return False


def structured_collect_record_fields(node, path_parts, max_depth, max_fields, current_depth=0):
    fields = []
    record_child_paths = structured_record_child_paths(node, path_parts, max_depth)

    def walk(value, current_path, depth):
        if len(fields) >= max_fields:
            return
        if depth > max_depth:
            return
        if current_path != path_parts and structured_should_skip_subtree(current_path, record_child_paths):
            return

        if structured_is_scalar(value):
            text = structured_scalar_text(value)
            if text:
                fields.append((current_path, structured_path_label(current_path), text))
            return

        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, current_path + [str(key)], depth + 1)
                if len(fields) >= max_fields:
                    break
        elif isinstance(value, list):
            for index, child in enumerate(value):
                walk(child, current_path + [f"[{index}]"], depth + 1)
                if len(fields) >= max_fields:
                    break

    walk(node, path_parts, current_depth)
    return fields


def structured_add_synthesised_fields(fields):
    name_fields = []
    for path, label, value in fields:
        metadata = resolve_table_header_metadata(label)
        name_fields.append((label, value, metadata))

    for name in assembled_names_from_fields(name_fields):
        fields.append(([], "full name", name["matched_text"]))

    return fields


def structured_fields_to_line(record_type, record_path, fields, source_name=""):
    fields = structured_add_synthesised_fields(list(fields))
    labelled = []
    seen = set()
    for _path, label, value in fields:
        label = safe_cell(label) or "value"
        value = safe_cell(value)
        if not value:
            continue
        key = (label, value)
        if key in seen:
            continue
        seen.add(key)
        labelled.append(f"{label}: {value}")

    if not labelled:
        return ""

    source_suffix = f" {source_name}" if source_name else ""
    return f"{record_type}{source_suffix} {record_path} | " + " | ".join(labelled)



def structured_emit_self_record_if_evidence(node, path_parts, record_type, max_depth, max_records, max_fields, parent_context, records):
    """Emit a self-record for a container that also has child records.

    This is the key safeguard for structures such as:

        student
          STUDENT_ID
          FIRST_NAME
          LAST_NAME
          BIRTH_DATE
          email[]
          address[]
          programme[]

    The parent node is both a person-bearing record and a container for repeated
    child evidence records. We emit the parent node's own scalar evidence while
    excluding nested record children, then recurse into those child records
    separately. This avoids two failure modes:

    1. Losing parent-level demographics/identifiers because the node was treated
       only as a container.
    2. Squashing sibling/repeated child records into one unsafe mega-record.
    """
    if max_records and len(records) >= max_records:
        return False

    self_fields = structured_collect_record_fields(
        node,
        path_parts,
        max_depth=max_depth,
        max_fields=max_fields,
    )

    if structured_evidence_score(self_fields)[2] < 2:
        return False

    fields = list(parent_context) + self_fields
    if len(fields) > max_fields:
        fields = fields[:max_fields]

    line = structured_fields_to_line(record_type, structured_path_text(path_parts), fields)
    if not line:
        return False

    records.append(line)
    return True

def structured_emit_records(node, path_parts, record_type, max_depth, max_records, max_fields, parent_context=None, records=None):
    if parent_context is None:
        parent_context = []
    if records is None:
        records = []
    if max_records and len(records) >= max_records:
        return records
    if structured_is_scalar(node):
        return records

    role = structured_node_role(path_parts, node, max_depth)
    child_record_paths = structured_record_child_paths(node, path_parts, max_depth)
    child_record_count = len(child_record_paths)
    is_record = structured_is_record_candidate(node, path_parts, max_depth)

    current_context = parent_context + structured_safe_context_fields(node, path_parts)

    # If this node contains several person-like child records, the parent is a
    # context/container scope. Do not squash the children together. However, the
    # parent may also have its own direct person evidence, for example a student
    # demographics node that contains STUDENT_ID/FIRST_NAME/LAST_NAME/BIRTH_DATE
    # as well as repeated child address/email/programme records. Emit a bounded
    # self-record for the parent first, excluding nested record children, then
    # recurse into the child records separately.
    if child_record_count > 1:
        structured_emit_self_record_if_evidence(
            node,
            path_parts,
            record_type,
            max_depth=max_depth,
            max_records=max_records,
            max_fields=max_fields,
            parent_context=parent_context,
            records=records,
        )
        if max_records and len(records) >= max_records:
            return records
        for child_path, child in structured_child_containers(node, path_parts):
            structured_emit_records(child, child_path, record_type, max_depth, max_records, max_fields, current_context, records)
            if max_records and len(records) >= max_records:
                break
        return records

    # If the only person-like evidence is deeper inside one child, let the lower
    # node become the record and inherit safe context from here.
    if child_record_count == 1 and not role["has_direct_identity"] and role["direct_score"] < 2 and not structured_is_collection_item(path_parts):
        for child_path, child in structured_child_containers(node, path_parts):
            structured_emit_records(child, child_path, record_type, max_depth, max_records, max_fields, current_context, records)
            if max_records and len(records) >= max_records:
                break
        return records

    if is_record:
        fields = []
        fields.extend(parent_context)
        fields.extend(structured_collect_record_fields(node, path_parts, max_depth, max_fields))
        if len(fields) > max_fields:
            fields = fields[:max_fields]
        line = structured_fields_to_line(record_type, structured_path_text(path_parts), fields)
        if line:
            records.append(line)
        return records

    for child_path, child in structured_child_containers(node, path_parts):
        structured_emit_records(child, child_path, record_type, max_depth, max_records, max_fields, current_context, records)
        if max_records and len(records) >= max_records:
            break

    # Fallback: if nothing below emitted but this node itself has linked-looking
    # evidence and no unsafe sibling squashing is detected, emit this scope.
    if not records and role["descendant_score"] >= 2:
        fields = parent_context + structured_collect_record_fields(node, path_parts, max_depth, max_fields)
        line = structured_fields_to_line(record_type, structured_path_text(path_parts), fields)
        if line:
            records.append(line)

    return records


def structured_data_to_labelled_text(data, record_type, max_depth=12, max_records=0, max_fields=400):
    records = structured_emit_records(
        data,
        ["$"],
        record_type,
        max_depth=max_depth,
        max_records=max_records,
        max_fields=max_fields,
    )
    if not records:
        fields = structured_all_scalar_fields(data, ["$"], max_depth)
        strong, supporting, score = structured_evidence_score(fields)
        if score >= 2:
            line = structured_fields_to_line(record_type, "$", fields[:max_fields])
            if line:
                records.append(line)
    return "\n".join(records)


def extract_json_structured(path, max_depth=12, max_records=0, max_fields=400):
    raw_text, status = read_text_sample(str(path), 0)
    if status == "binary_like_content":
        return ExtractedText("", status)
    data = json.loads(raw_text)
    text = structured_data_to_labelled_text(data, "JSON_RECORD", max_depth, max_records, max_fields)
    if text:
        return ExtractedText(text, "json_structured_records")
    return ExtractedText(raw_text, "plain_json_text")


def extract_jsonl_structured(path, max_depth=12, max_records=0, max_fields=400):
    raw_text, status = read_text_sample(str(path), 0)
    if status == "binary_like_content":
        return ExtractedText("", status)

    records = []
    parsed_any = False
    for line_number, line in enumerate(io.StringIO(raw_text), start=1):
        if max_records and len(records) >= max_records:
            break
        stripped = line.strip()
        if not stripped:
            continue
        try:
            data = json.loads(stripped)
        except json.JSONDecodeError:
            continue
        parsed_any = True
        emitted = structured_emit_records(
            data,
            ["$", f"[{line_number}]"],
            "JSONL_RECORD",
            max_depth=max_depth,
            max_records=max(0, max_records - len(records)) if max_records else 0,
            max_fields=max_fields,
        )
        records.extend(emitted)
        if not emitted:
            fields = structured_all_scalar_fields(data, ["$", f"[{line_number}]"], max_depth)
            if structured_evidence_score(fields)[2] >= 2:
                fallback = structured_fields_to_line("JSONL_RECORD", f"$[{line_number}]", fields[:max_fields])
                if fallback:
                    records.append(fallback)

    if records:
        return ExtractedText("\n".join(records), "jsonl_structured_records")
    if parsed_any:
        return ExtractedText(raw_text, "plain_jsonl_text")
    return ExtractedText(raw_text, "plain_text")


def xml_element_to_structured_data(element):
    data = {}

    for key, value in element.attrib.items():
        text = safe_cell(value)
        if text:
            data[f"@{key}"] = text

    text = safe_cell(element.text)
    children = list(element)
    if text and children:
        data["#text"] = text
    elif text and not children:
        return text

    grouped = defaultdict(list)
    for child in children:
        grouped[child.tag].append(xml_element_to_structured_data(child))

    for tag, values in grouped.items():
        data[tag] = values[0] if len(values) == 1 else values

    return data


def extract_xml_structured(path, max_depth=12, max_records=0, max_fields=400, progress_callback=None):
    import xml.etree.ElementTree as ET

    try:
        size_bytes = os.path.getsize(path)
    except OSError:
        size_bytes = 0
    if size_bytes >= XML_STREAMING_STRUCTURED_BYTES:
        streamed = extract_xml_structured_streaming(
            path,
            max_depth=max_depth,
            max_records=max_records,
            max_fields=max_fields,
            progress_callback=progress_callback,
        )
        if streamed is not None:
            return streamed

    if progress_callback:
        progress_callback(28.0, "reading XML")
    raw_text, status = read_text_sample(str(path), 0)
    if status == "binary_like_content":
        return ExtractedText("", status)
    if progress_callback:
        progress_callback(40.0, "parsing XML")
    root = ET.fromstring(raw_text)
    data = {root.tag: xml_element_to_structured_data(root)}
    if progress_callback:
        progress_callback(52.0, "emitting structured XML records")
    text = structured_data_to_labelled_text(data, "XML_RECORD", max_depth, max_records, max_fields)
    if progress_callback:
        progress_callback(60.0, "XML extraction complete")
    if text:
        return ExtractedText(text, "xml_structured_records")
    return ExtractedText(raw_text, "plain_xml_text")


def extract_xml_structured_streaming(path, max_depth=12, max_records=0, max_fields=400, progress_callback=None):
    import xml.etree.ElementTree as ET

    records = []
    emitted_count = 0
    stack = []
    root_tag = None
    parsed_any = False
    total_size = 0
    try:
        total_size = max(1, os.path.getsize(path))
    except OSError:
        total_size = 0
    last_progress_ratio = -1.0
    try:
        with open(path, "rb") as handle:
            for event, element in ET.iterparse(handle, events=("start", "end")):
                if progress_callback and total_size > 0:
                    try:
                        ratio = min(1.0, max(0.0, float(handle.tell()) / float(total_size)))
                    except Exception:
                        ratio = 0.0
                    if ratio - last_progress_ratio >= 0.02:
                        progress_callback(25.0 + (27.0 * ratio), "streaming XML records")
                        last_progress_ratio = ratio

                if event == "start":
                    stack.append(element.tag)
                    if root_tag is None:
                        root_tag = element.tag
                    continue

                parsed_any = True
                is_direct_root_child = root_tag is not None and len(stack) == 2
                if is_direct_root_child:
                    data = {root_tag: {element.tag: xml_element_to_structured_data(element)}}
                    remaining = max(0, max_records - emitted_count) if max_records else 0
                    text = structured_data_to_labelled_text(
                        data,
                        "XML_RECORD",
                        max_depth=max_depth,
                        max_records=remaining,
                        max_fields=max_fields,
                    )
                    if text:
                        records.append(text)
                        emitted_count += text.count("\n") + 1
                        if max_records and emitted_count >= max_records:
                            element.clear()
                            if progress_callback:
                                progress_callback(60.0, "XML extraction complete")
                            return ExtractedText("\n".join(records), "xml_streaming_structured_records")
                    element.clear()
                if stack:
                    stack.pop()
    except ET.ParseError:
        return None

    if progress_callback:
        progress_callback(60.0, "XML extraction complete")
    if records:
        return ExtractedText("\n".join(records), "xml_streaming_structured_records")
    if parsed_any:
        return None
    return ExtractedText("", "plain_xml_text")


def extract_yaml_structured(path, max_depth=12, max_records=0, max_fields=400):
    raw_text, status = read_text_sample(str(path), 0)
    if status == "binary_like_content":
        return ExtractedText("", status)
    try:
        import yaml
    except ImportError as exc:
        raise ValueError("PyYAML is not installed; cannot parse YAML structurally.") from exc
    data = yaml.safe_load(raw_text)
    text = structured_data_to_labelled_text(data, "YAML_RECORD", max_depth, max_records, max_fields)
    if text:
        return ExtractedText(text, "yaml_structured_records")
    return ExtractedText(raw_text, "plain_yaml_text")


def extract_delimited_table(path, delimiter=None):
    raw, status = read_text_sample(str(path), 0)
    if status == "binary_like_content":
        return ExtractedText("", status)
    sample = raw[:8192]
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=TABLE_SNIFF_DELIMITERS)
        except csv.Error:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
            dialect = csv.excel_tab if delimiter == "\t" else csv.excel
    else:
        dialect = csv.excel_tab if delimiter == "\t" else csv.excel

    rows = csv.reader(io.StringIO(raw, newline=""), dialect)
    try:
        headers = next(rows)
    except StopIteration:
        return ExtractedText("", "delimited_empty")
    text = table_rows_to_labelled_text(headers, rows, path.name)
    if text:
        return ExtractedText(text, "labelled_delimited_table")
    return ExtractedText(raw, "plain_text")

def extract_text_like(
    path,
    max_text_chars,
    structured_records=True,
    structured_max_depth=12,
    structured_max_records=0,
    structured_max_record_fields=400,
    progress_callback=None,
):
    extension = path.suffix.lower()

    if structured_records and extension == ".json":
        try:
            return extract_json_structured(
                path,
                max_depth=structured_max_depth,
                max_records=structured_max_records,
                max_fields=structured_max_record_fields,
            )
        except Exception:
            pass

    if structured_records and extension == ".jsonl":
        try:
            return extract_jsonl_structured(
                path,
                max_depth=structured_max_depth,
                max_records=structured_max_records,
                max_fields=structured_max_record_fields,
            )
        except Exception:
            pass

    if structured_records and extension == ".xml":
        try:
            return extract_xml_structured(
                path,
                max_depth=structured_max_depth,
                max_records=structured_max_records,
                max_fields=structured_max_record_fields,
                progress_callback=progress_callback,
            )
        except Exception:
            pass

    if structured_records and extension in {".yaml", ".yml"}:
        try:
            return extract_yaml_structured(
                path,
                max_depth=structured_max_depth,
                max_records=structured_max_records,
                max_fields=structured_max_record_fields,
            )
        except Exception:
            pass

    if extension in {".csv", ".tsv"}:
        delimiter = "	" if extension == ".tsv" else None
        extracted = extract_delimited_table(path, delimiter=delimiter)
        if extracted.extraction_method == "labelled_delimited_table":
            return extracted

    text, status = read_text_sample(str(path), max_text_chars)
    if status == "binary_like_content":
        return ExtractedText("", status)

    if extension in {".html", ".htm", ".xml"}:
        from bs4 import BeautifulSoup

        parser = "xml" if extension == ".xml" else "lxml"
        soup = BeautifulSoup(text, parser)
        method = "truncated_beautifulsoup" if status == "truncated" else "beautifulsoup"
        return ExtractedText(soup.get_text("\n"), method)

    embedded_tables_text = labelled_embedded_delimited_blocks(text, path.name)
    if embedded_tables_text:
        method = "truncated_plain_text_with_detected_delimited_blocks" if status == "truncated" else "plain_text_with_detected_delimited_blocks"
        return ExtractedText(embedded_tables_text, method)

    method = "truncated_plain_text" if status == "truncated" else "plain_text"
    return ExtractedText(text, method)


def dataframes_to_labelled_text(frames, method):
    parts = []
    for sheet_name, df in frames.items():
        text = table_rows_to_labelled_text(df.columns.tolist(), df.fillna("").itertuples(index=False, name=None), sheet_name)
        if text:
            parts.append(text)
        else:
            parts.append(f"\n--- SHEET: {sheet_name} ---\n")
            parts.append(df.fillna("").to_csv(index=False))
    return ExtractedText("\n".join(parts), method)


def extract_openpyxl_workbook(path, method="openpyxl_read_only_labelled_tables"):
    from openpyxl import load_workbook

    parts = []
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=r".*Data Validation extension is not supported and will be removed.*",
            category=UserWarning,
            module=r"openpyxl\..*",
        )
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = next(iterator)
            except StopIteration:
                continue
            text = table_rows_to_labelled_text(headers, iterator, sheet.title)
            if text:
                parts.append(text)
        return ExtractedText("\n".join(parts), method)
    finally:
        workbook.close()


def extract_pandas_excel(path, engine):
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None, dtype=str, engine=engine)
    return dataframes_to_labelled_text(sheets, f"pandas_excel_{engine}_labelled_tables")


def extract_html_tables_as_excel(path):
    import pandas as pd

    tables = pd.read_html(path, flavor="lxml")
    if not tables:
        raise ValueError("No HTML tables found.")
    frames = {f"html_table_{index}": df for index, df in enumerate(tables, start=1)}
    return dataframes_to_labelled_text(frames, "pandas_html_xls_fallback_labelled_tables")


def extract_markup_text_as_excel(path):
    text, status = read_text_sample(str(path), 0)
    if not text.strip():
        raise ValueError("No readable text found.")

    from bs4 import BeautifulSoup

    parser = "xml" if text.lstrip().lower().startswith("<?xml") else "lxml"
    soup = BeautifulSoup(text, parser)
    method = "markup_text_xls_fallback" if status == "scanned" else f"{status}_markup_text_xls_fallback"
    return ExtractedText(soup.get_text("\n"), method)


def extract_plain_text_as_excel(path):
    text, status = read_text_sample(str(path), 0)
    if not text.strip():
        raise ValueError("No readable text found.")
    method = "plain_text_xls_fallback" if status == "scanned" else f"{status}_plain_text_xls_fallback"
    return ExtractedText(text, method)


def extract_delimited_text_as_excel(path):
    raw_text, status = read_text_sample(str(path), 0)
    if status == "binary_like_content":
        raise ValueError("Binary-like content is not a delimited text fallback.")
    if not raw_text.strip():
        raise ValueError("No delimited text found.")

    sample = raw_text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=TABLE_SNIFF_DELIMITERS)
    except csv.Error as exc:
        raise ValueError("Could not detect a delimited text dialect.") from exc

    rows = csv.reader(io.StringIO(raw_text, newline=""), dialect)
    try:
        headers = next(rows)
    except StopIteration as exc:
        raise ValueError("Delimited text fallback has no rows.") from exc

    text = table_rows_to_labelled_text(headers, rows, path.name)
    if text:
        return ExtractedText(text, "labelled_delimited_table_xls_fallback")
    return ExtractedText(raw_text, "plain_delimited_text_xls_fallback")


def try_excel_readers(path, readers):
    errors = []
    seen = set()
    for method_name, reader in readers:
        if method_name in seen:
            continue
        seen.add(method_name)
        try:
            extracted = reader()
        except Exception as exc:
            errors.append(f"{method_name}: {type(exc).__name__}: {exc}")
            continue
        if extracted.text.strip():
            return extracted
        errors.append(f"{method_name}: no extractable text")
    raise ValueError("Unable to read spreadsheet with available readers. " + " | ".join(errors))


def extract_excel(path):
    extension = path.suffix.lower()

    if extension in {".xlsx", ".xlsm"}:
        return try_excel_readers(
            path,
            [
                ("openpyxl", lambda: extract_openpyxl_workbook(path)),
                ("html_tables", lambda: extract_html_tables_as_excel(path)),
                ("delimited_text", lambda: extract_delimited_text_as_excel(path)),
                ("markup_text", lambda: extract_markup_text_as_excel(path)),
                ("plain_text", lambda: extract_plain_text_as_excel(path)),
            ],
        )

    if extension == ".xlsb":
        engine = "pyxlsb"
    elif extension == ".ods":
        engine = "odf"
    elif extension == ".xls":
        engine = "xlrd"
    else:
        engine = "openpyxl"

    return try_excel_readers(
        path,
        [
            (f"pandas_excel_{engine}", lambda: extract_pandas_excel(path, engine)),
            ("openpyxl_mislabelled_ooxml", lambda: extract_openpyxl_workbook(path, "openpyxl_mislabelled_ooxml_fallback")),
            ("html_tables", lambda: extract_html_tables_as_excel(path)),
            ("delimited_text", lambda: extract_delimited_text_as_excel(path)),
            ("markup_text", lambda: extract_markup_text_as_excel(path)),
            ("plain_text", lambda: extract_plain_text_as_excel(path)),
        ],
    )

def extract_docx(path):
    from docx import Document

    doc = Document(path)
    parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"\n--- TABLE: {table_index} ---")
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))

    return ExtractedText("\n".join(parts), "python_docx")


def extract_pdf(path):
    from pypdf import PdfReader

    parts = []
    reader = PdfReader(path)
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"\n--- PDF PAGE: {page_number} ---\n{text}")
    return ExtractedText("\n".join(parts), "pypdf_embedded_text")


def extract_image_ocr(path, ocr_images):
    if not ocr_images:
        return ExtractedText("", "ocr_disabled")

    from PIL import Image
    import pytesseract

    with Image.open(path) as image:
        text = pytesseract.image_to_string(image)
    return ExtractedText(text, "pytesseract_ocr")


def extract_rtf(path):
    from striprtf.striprtf import rtf_to_text

    raw = path.read_text(encoding="utf-8", errors="ignore")
    return ExtractedText(rtf_to_text(raw), "striprtf")


def extract_msg(path):
    import extract_msg

    msg = extract_msg.Message(str(path))
    parts = [
        f"Subject: {msg.subject or ''}",
        f"Sender: {msg.sender or ''}",
        f"To: {msg.to or ''}",
        f"CC: {msg.cc or ''}",
        "",
        msg.body or "",
    ]
    return ExtractedText("\n".join(parts), "extract_msg")


def extract_eml(path):
    import email
    from email import policy

    with path.open("rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)

    parts = [
        f"Subject: {msg.get('subject', '')}",
        f"From: {msg.get('from', '')}",
        f"To: {msg.get('to', '')}",
        f"CC: {msg.get('cc', '')}",
        "",
    ]

    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                parts.append(part.get_content())
    elif msg.get_content_type() == "text/plain":
        parts.append(msg.get_content())

    return ExtractedText("\n".join(parts), "email_parser")


def extract_text_for_scan(
    path,
    extension,
    max_text_chars,
    ocr_images,
    structured_records=True,
    structured_max_depth=12,
    structured_max_records=0,
    structured_max_record_fields=400,
    progress_callback=None,
):
    path_obj = Path(path)

    if extension in TEXT_EXTENSIONS:
        extracted = extract_text_like(
            path_obj,
            max_text_chars,
            structured_records=structured_records,
            structured_max_depth=structured_max_depth,
            structured_max_records=structured_max_records,
            structured_max_record_fields=structured_max_record_fields,
            progress_callback=progress_callback,
        )
    elif extension in EXCEL_EXTENSIONS:
        extracted = extract_excel(path_obj)
    elif extension in WORD_EXTENSIONS:
        extracted = extract_docx(path_obj)
    elif extension in PDF_EXTENSIONS:
        extracted = extract_pdf(path_obj)
    elif extension in IMAGE_EXTENSIONS:
        extracted = extract_image_ocr(path_obj, ocr_images)
    elif extension in RTF_EXTENSIONS:
        extracted = extract_rtf(path_obj)
    elif extension == ".msg":
        extracted = extract_msg(path_obj)
    elif extension == ".eml":
        extracted = extract_eml(path_obj)
    else:
        return "", "unsupported_extension"

    if extracted.extraction_method == "binary_like_content":
        return "", "binary_like_content"
    return limit_extracted_text(extracted, max_text_chars)


# =============================================================================

# Merged from scanning(3).py

# =============================================================================

"""Target discovery, source-unit scanning, evidence scoring, and CSV summary helpers."""



def sync_runtime_state():
    return None


def excluded_path_config(root, exclude_groups):
    excluded_paths = set()
    excluded_names = set()
    excluded_suffixes = set()
    for group in exclude_groups or []:
        values = group if isinstance(group, (list, tuple)) else [group]
        for value in values:
            for item in str(value).split(","):
                item = item.strip()
                if not item:
                    continue
                expanded = os.path.expanduser(item)
                portable_path = expanded.replace("\\", "/")
                if os.path.isabs(expanded):
                    path = expanded
                    excluded_paths.add(os.path.abspath(path))
                elif "/" in portable_path:
                    native_path = os.path.normpath(portable_path.replace("/", os.sep))
                    excluded_paths.add(os.path.abspath(os.path.join(root, native_path)))
                    suffix_parts = tuple(
                        part.casefold()
                        for part in native_path.replace(os.sep, "/").split("/")
                        if part not in {"", ".", ".."}
                    )
                    if suffix_parts:
                        excluded_suffixes.add(suffix_parts)
                else:
                    excluded_names.add(item.casefold())
    return excluded_paths, excluded_names, excluded_suffixes


def is_excluded_path(path, root, excluded_paths, excluded_names, excluded_suffixes):
    absolute_path = os.path.abspath(path)
    if os.path.basename(absolute_path).casefold() in excluded_names:
        return True
    if any(absolute_path == folder or is_inside(absolute_path, folder) for folder in excluded_paths):
        return True

    try:
        relative_path = os.path.relpath(absolute_path, root)
    except ValueError:
        return False
    relative_parts = tuple(part.casefold() for part in Path(relative_path).parts)
    return any(
        len(relative_parts) >= len(suffix) and relative_parts[-len(suffix):] == suffix
        for suffix in excluded_suffixes
    )


def iter_targets(root, output_dir, output_files, include_hidden, exclude_groups=None):
    stack = [root]
    output_dir = os.path.abspath(output_dir)
    output_files = {os.path.abspath(path) for path in output_files}
    excluded_paths, excluded_names, excluded_suffixes = excluded_path_config(root, exclude_groups)
    skip_output_dir = output_dir != os.path.abspath(root) and is_inside(output_dir, root)

    while stack:
        dirpath = stack.pop()

        try:
            with os.scandir(dirpath) as entries:
                visible_entries = list(entries)
                if not include_hidden:
                    visible_entries = [entry for entry in visible_entries if not entry.name.startswith(".")]
                visible_entries = [entry for entry in visible_entries if os.path.abspath(entry.path) not in output_files]

                if not visible_entries and not is_inside(dirpath, output_dir):
                    yield "empty_folder", dirpath

                for entry in visible_entries:
                    path = entry.path

                    if skip_output_dir and is_inside(path, output_dir):
                        continue
                    if is_excluded_path(path, root, excluded_paths, excluded_names, excluded_suffixes):
                        continue

                    try:
                        if entry.is_dir(follow_symlinks=False):
                            stack.append(path)
                        elif entry.is_file(follow_symlinks=False):
                            yield "file", path
                    except OSError:
                        yield "file", path
        except OSError:
            continue


def target_priority_key(target, priority):
    target_type, path = target
    normalized_path = os.path.normcase(os.path.abspath(path)).casefold()

    if priority == "alphabetical":
        return (normalized_path,)
    if priority == "type":
        extension = Path(path).suffix.casefold() if target_type == "file" else ""
        return (target_type != "file", extension, normalized_path)
    if priority == "size":
        try:
            size_bytes = os.path.getsize(path) if target_type == "file" else 0
        except OSError:
            size_bytes = -1
        return (-size_bytes, normalized_path)
    return (0,)


def prioritize_targets(targets, priority):
    if priority == "discovery":
        return targets
    return sorted(targets, key=lambda target: target_priority_key(target, priority))


def is_inside(path, folder):
    try:
        normalized_folder = os.path.abspath(folder)
        return os.path.commonpath([os.path.abspath(path), normalized_folder]) == normalized_folder
    except ValueError:
        return False
